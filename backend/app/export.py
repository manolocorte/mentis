"""Render a whitepaper (markdown) to a presentable academic document — PDF (default)
or Word (.docx). Pure-Python (markdown + xhtml2pdf + python-docx): no LaTeX, no headless
browser, so it runs the same locally and on Lambda.
"""
from __future__ import annotations

import io
import re

import markdown as md
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from xhtml2pdf import pisa

# --- PDF (academic single-column) ---
_CSS = """
@page {
    size: a4;
    margin: 2.5cm 2.3cm 2.6cm 2.3cm;
    @frame footer { -pdf-frame-content: footerContent; bottom: 1.2cm; margin-left: 2.3cm; margin-right: 2.3cm; height: 1cm; }
}
body { font-family: "Times New Roman", Times, serif; font-size: 11pt; line-height: 1.55; text-align: justify; color: #111; }
.doc-title { font-size: 18pt; font-weight: bold; text-align: center; margin: 0 0 4pt 0; line-height: 1.25; }
.rule { border-bottom: 0.75pt solid #444; margin: 6pt 0 12pt 0; }
h1 { font-size: 18pt; font-weight: bold; text-align: center; margin: 0 0 12pt 0; }
h2 { font-size: 13pt; font-weight: bold; margin: 14pt 0 5pt 0; }
h3 { font-size: 11.5pt; font-weight: bold; margin: 11pt 0 4pt 0; }
p { margin: 0 0 7pt 0; }
a { color: #1a4f8b; text-decoration: none; }
ol, ul { margin: 3pt 0 7pt 18pt; }
li { margin: 0 0 3pt 0; }
code { font-family: "Courier New", monospace; font-size: 10pt; }
"""

_FOOTER = (
    '<div id="footerContent" style="text-align:center; font-size:8pt; color:#888;">'
    'Mentis · page <pdf:pagenumber/></div>'
)


def _extract_title(text: str) -> str | None:
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("# "):
            return s[2:].strip()
    return None


def build_pdf(markdown_text: str, title: str | None = None) -> bytes:
    body_html = md.markdown(markdown_text or "", extensions=["extra", "sane_lists"])
    has_h1 = "<h1" in body_html
    head = ""
    if title:
        head = f'<div class="doc-title">{title}</div><div class="rule"></div>'
    elif not has_h1:
        head = '<div class="doc-title">Mentis Whitepaper</div><div class="rule"></div>'
    html = (
        f'<html><head><meta charset="utf-8"><style>{_CSS}</style></head>'
        f"<body>{_FOOTER}{head}{body_html}</body></html>"
    )
    out = io.BytesIO()
    if pisa.CreatePDF(src=html, dest=out, encoding="utf-8").err:
        raise RuntimeError("PDF generation failed")
    return out.getvalue()


# --- Word (.docx, academic) ---
_INLINE = [
    (re.compile(r"\*\*(.+?)\*\*"), r"\1"),
    (re.compile(r"(?<!\*)\*(?!\*)(.+?)\*"), r"\1"),
    (re.compile(r"`(.+?)`"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)"), r"\1 (\2)"),
]


def _clean_inline(text: str) -> str:
    for pat, repl in _INLINE:
        text = pat.sub(repl, text)
    return text.strip()


def _heading(doc: Document, text: str, size: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(_clean_inline(text))
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)


def build_docx(markdown_text: str, title: str | None = None) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc_title = title or _extract_title(markdown_text) or "Mentis Whitepaper"
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(doc_title)
    tr.bold = True
    tr.font.size = Pt(16)

    for raw in (markdown_text or "").splitlines():
        line = raw.rstrip()
        s = line.strip()
        if not s:
            continue
        if s.startswith("# "):
            if s[2:].strip() == doc_title:
                continue
            _heading(doc, s[2:], 14)
        elif s.startswith("### "):
            _heading(doc, s[4:], 11)
        elif s.startswith("## "):
            _heading(doc, s[3:], 13)
        elif s.startswith(("- ", "* ")):
            doc.add_paragraph(_clean_inline(s[2:]), style="List Bullet")
        else:
            doc.add_paragraph(_clean_inline(s))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
